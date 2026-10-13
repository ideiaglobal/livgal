"""Lista de candidatos a «ediçom erudita» (crítica, anotada, com estudo/introduçom...) na BD livpoesia.

Uso: python3 candidatos_edicoes_eruditas.py [export.json]
Gera candidatos_edicoes_eruditas.xlsx e .csv nesta pasta. A CDU nom se usa.
"""
import csv, re, sys, unicodedata, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import sinais as S

HERE = Path(__file__).parent
EXPORT = sys.argv[1] if len(sys.argv) > 1 else HERE.parent / 'atividade-editorial' / 'Export_livpoesia_20261012104901.json'
CODES = [c for c, *_ in S.SIGNALS]
FLAG_TAGS = ['duplicados_poesia', 'coletivos', 'traduçom_para', 'poesia_infantil', 'poesia_escolar', 'exposicións', 'catálogos']


def norm(t):
    t = unicodedata.normalize('NFD', t.lower())
    return ''.join(c for c in t if unicodedata.category(c) != 'Mn')


def year(v):
    m = re.search(r'\d{4}', str(v))
    return int(m.group()) if m else None


def editor_is_author(r):
    """True se todos os «ed. lit.» dos autores secundários coincidem com o apelido da autoria principal."""
    authors = {norm(a.split(',')[0]) for a in S.vals(r, 'Autoria')}
    eds = set()
    for v in S.vals(r, 'Autores secundários (marc_map)'):
        for m in re.finditer(r'([A-ZÁÉÍÓÚÑ][^,.;]+),[^.;]*?,\s*e?d\.\s?lit', v):
            eds.add(norm(m.group(1).split()[-2] if len(m.group(1).split()) > 2 else m.group(1)))
    return bool(eds) and bool(authors) and all(any(e in a or a in e for a in authors) for e in eds)


def level(sc, found):
    if 'critica' in found or sc >= 5:
        return 'A'
    if sc >= 3:
        return 'B'
    return 'C' if sc >= 1 else '-'


def build():
    books, rels, people = S.load(EXPORT)
    rows = []
    for r in books:
        found = S.detect(r, rels)
        tags = S.vals(r, 'Etiqueta')
        tagged = S.TAG in tags
        pub = year((S.vals(r, 'Ano de publicação verificado') or S.vals(r, 'Ano de publicação') or [''])[0])
        # autor falecido antes da publicaçom (só pesa se houver outro sinal)
        dead = False
        for d in r['details']:
            if d['fieldName'] == 'Autoria' and isinstance(d['value'], dict):
                p = people.get(d['value']['id'])
                dy = year((S.vals(p, 'Data da morte') or [''])[0]) if p else None
                if dy and pub and dy < pub:
                    dead = True
        self_ed = editor_is_author(r)
        if self_ed and not any(f != 'Autores secundários (marc_map)' for f, _ in found.get('edicao', [])):
            found.pop('edicao', None)
        sc = S.score(found) + (2 if dead and S.score(found) > 0 else 0)
        if sc == 0 and not tagged:
            continue
        lv = level(sc, found)
        estado = ('já etiquetado' if lv != '-' else 'já etiquetado, sem sinais (rever)') if tagged else 'novo candidato'
        evid = ' ‖ '.join(f'[{c}] {f}: {t}' for c in CODES if c != 'prologo' for f, t in found.get(c, [])[:2])
        row = {
            'rec_ID': int(r['rec_ID']),
            'nível': lv, 'pontos': sc, 'estado': estado,
            'título': ' : '.join((S.vals(r, 'Título principal')[:1] + S.vals(r, 'Subtítulo')[:1])),
            'autoria': '; '.join(S.vals(r, 'Autoria')) or '(sem autoria principal)',
            'ano': pub, 'editora': '; '.join(dict.fromkeys(S.vals(r, 'Editora'))),
            'natureza': '; '.join(S.term(r, 'Natureza da edição')),
        }
        for c in CODES:
            row[c] = 'x' if c in found else ''
        row['autor falecido'] = 'x' if dead else ''
        row['editor = autor'] = 'x' if self_ed else ''
        row['outras etiquetas'] = ', '.join(t for t in FLAG_TAGS if t in tags)
        row['evidência'] = evid
        row['mençom de responsabilidade'] = max(S.vals(r, 'Autoria (marc_map)'), key=len, default='')
        row['decisom (S/N/?)'] = ''
        rows.append(row)
    order = {'A': 0, 'B': 1, 'C': 2, '-': 3}
    rows.sort(key=lambda x: (order[x['nível']], -x['pontos'], x['título'].lower()))
    return rows, len(books)


def write(rows, nbooks):
    cols = list(rows[0])
    with open(HERE / 'candidatos_edicoes_eruditas.csv', 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, cols)
        w.writeheader()
        w.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Candidatos'
    ws.append(cols)
    fills = {'A': 'C6E0B4', 'B': 'FFE699', 'C': 'F2F2F2', '-': 'F8CBAD'}
    for row in rows:
        ws.append([row[c] for c in cols])
        ws.cell(ws.max_row, 2).fill = PatternFill('solid', start_color=fills[row['nível']])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    widths = {'rec_ID': 8, 'nível': 6, 'pontos': 7, 'estado': 24, 'título': 50, 'autoria': 28, 'ano': 6, 'editora': 24,
              'natureza': 13, 'outras etiquetas': 20, 'evidência': 90, 'mençom de responsabilidade': 70, 'decisom (S/N/?)': 10}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 11)
    ws.freeze_panes = 'F2'
    ws.auto_filter.ref = ws.dimensions

    rs = wb.create_sheet('Resumo')
    cnt = collections.Counter((r['nível'], r['estado']) for r in rows)
    estados = ['já etiquetado', 'novo candidato', 'já etiquetado, sem sinais (rever)']
    rs.append(['nível'] + estados + ['total'])
    for lv in 'ABC-':
        v = [cnt[(lv, e)] for e in estados]
        rs.append([lv] + v + [sum(v)])
    rs.append(['total'] + [sum(cnt[(lv, e)] for lv in 'ABC-') for e in estados] + [len(rows)])
    rs.append([])
    rs.append(['sinal', 'descriçom', 'peso', 'livros', 'já etiquetados'])
    for c, d, w, _, _ in S.SIGNALS:
        rs.append([c, d, w, sum(1 for r in rows if r[c]), sum(1 for r in rows if r[c] and r['estado'].startswith('já'))])
    rs.append([])
    for line in [
        f'Universo: {nbooks} registos «Livro»; a CDU nom se usa.',
        f'«já etiquetado» = registo com a etiqueta «{S.TAG}» no export.',
        'Nível A: «ediçom crítica» literal ou 5+ pontos. B: 3–4 pontos. C: 1–2 pontos (sinal isolado e fraco).',
        'Pontos = soma dos pesos dos sinais presentes; +2 se o autor principal faleceu antes do ano de publicaçom.',
        '«prologo» (peso 0) é só informativo: sozinho nom gera candidato.',
        '«editor = autor»: o «ed. lit.» coincide com o autor principal (provável auto-ediçom, nom ediçom erudita).',
        'Sinais procurados nas mençons de responsabilidade, autores secundários, notas, subtítulo e funçons das relaçons.',
    ]:
        rs.append([line])
    for i, w in enumerate([14, 46, 16, 16, 30], 1):
        rs.column_dimensions[get_column_letter(i)].width = w
    for c in rs[1]:
        c.font = Font(bold=True)
    wb.save(HERE / 'candidatos_edicoes_eruditas.xlsx')
    return cnt


if __name__ == '__main__':
    rows, n = build()
    cnt = write(rows, n)
    for k in sorted(cnt):
        print(k, cnt[k])
    print('total', len(rows))
